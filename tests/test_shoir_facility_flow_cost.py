"""Tests for facility weighted flow-distance, handling cost, and provenance."""

import math

import pandas as pd
import pytest
from openpyxl import load_workbook
from io import BytesIO

import shoir_facility_layout as facility


@pytest.fixture
def flows():
    return pd.DataFrame([
        {"flow_id": "F001", "from_id": "D1", "to_id": "D2", "loads_day": 120, "distance_m": 22.0},
        {"flow_id": "F002", "from_id": "D2", "to_id": "D3", "loads_day": 150, "distance_m": 30.0},
    ])


def test_weighted_flow_distance_and_user_supplied_cost_rate(flows):
    # 120*22 + 150*30 = 7,140 load-metres/day; at 0.10 SAR/load-m = SAR 714/day.
    result = facility._flow_distance_metrics(flows, 0.10)
    assert result["status"] == "COMPLETED"
    assert result["connection_count"] == 2
    assert result["weighted_flow_distance_load_m_per_day"] == pytest.approx(7140.0)
    assert result["cost_rate_sar_per_load_meter"] == pytest.approx(0.10)
    assert result["estimated_daily_handling_cost_sar"] == pytest.approx(714.0)
    assert "uniform SAR per load-metre rate" in result["warnings"][0]


def test_weighted_flow_distance_is_available_without_inventing_cost(flows):
    result = facility._flow_distance_metrics(flows)
    assert result["weighted_flow_distance_load_m_per_day"] == pytest.approx(7140.0)
    assert result["estimated_daily_handling_cost_sar"] is None
    assert any("not calculated" in warning for warning in result["warnings"])


def test_empty_flow_register_is_not_presented_as_a_cost_estimate():
    result = facility._flow_distance_metrics(pd.DataFrame())
    assert result["status"] == "NO_DATA"
    assert result["weighted_flow_distance_load_m_per_day"] == 0
    assert result["estimated_daily_handling_cost_sar"] is None


@pytest.mark.parametrize(
    "bad_row",
    [
        {"flow_id": "F1", "loads_day": -1, "distance_m": 20},
        {"flow_id": "F1", "loads_day": 10, "distance_m": -1},
        {"flow_id": "F1", "loads_day": float("nan"), "distance_m": 20},
        {"flow_id": "F1", "loads_day": 10, "distance_m": float("inf")},
        {"flow_id": "F1", "loads_day": "bad", "distance_m": 20},
    ],
)
def test_flow_metrics_reject_invalid_measurements(bad_row):
    with pytest.raises(ValueError):
        facility._flow_distance_metrics(pd.DataFrame([bad_row]))


@pytest.mark.parametrize("rate", [0, -0.1, float("nan"), float("inf")])
def test_flow_metrics_reject_invalid_supplied_cost_rates(flows, rate):
    with pytest.raises(ValueError, match="Handling cost rate"):
        facility._flow_distance_metrics(flows, rate)


def test_workspace_data_origin_is_scoped_and_round_trips(tmp_path, monkeypatch):
    monkeypatch.setattr(facility, "DB_PATH", str(tmp_path / "facility-origin.db"))
    assert facility._get_data_origin("alice", "line-a") == "EXISTING_DATA_ORIGIN_UNVERIFIED"
    facility._set_data_origin("alice", "line-a", "SYNTHETIC_STARTER_DATA")
    facility._set_data_origin("alice", "line-b", "IMPORTED_OR_UPSERTED")
    facility._set_data_origin("bob", "line-a", "USER_EDITED_OR_MIXED")
    assert facility._get_data_origin("alice", "line-a") == "SYNTHETIC_STARTER_DATA"
    assert facility._get_data_origin("alice", "line-b") == "IMPORTED_OR_UPSERTED"
    assert facility._get_data_origin("bob", "line-a") == "USER_EDITED_OR_MIXED"


def test_export_reopens_with_flow_metrics_and_provenance_sheets(flows):
    depts = pd.DataFrame([
        {"id": "D1", "name": "Receiving", "x": 10, "y": 20, "area_sqm": 100},
        {"id": "D2", "name": "Assembly", "x": 50, "y": 60, "area_sqm": 200},
        {"id": "D3", "name": "Shipping", "x": 80, "y": 30, "area_sqm": 300},
    ])
    slp = facility._slp_frame(depts, flows)
    payload = facility._export_workbook(
        depts, flows, slp, "IMPORTED_OR_UPSERTED", 0.10
    )
    assert payload and len(payload) > 1000
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    assert workbook.sheetnames == [
        "Departments", "From_To_Flows", "SLP_Relationships", "SLP_Matrix",
        "Flow_Metrics", "Provenance",
    ]
    metric_rows = list(workbook["Flow_Metrics"].iter_rows(values_only=True))
    assert ("Metric", "Value") in metric_rows
    assert any(row[0] == "Estimated handling cost (SAR/day)" and float(row[1]) == pytest.approx(714.0) for row in metric_rows[1:])
    provenance_rows = list(workbook["Provenance"].iter_rows(values_only=True))
    assert any(row[0] == "Facility data-origin status" and row[1] == "IMPORTED_OR_UPSERTED" for row in provenance_rows[1:])
    workbook.close()
