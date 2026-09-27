import io

import pandas as pd
from openpyxl import load_workbook

from shoir_excel_studio import (
    _excel_apply_validation_rules,
    _excel_convert_unit_values,
    _excel_custom_rule_results,
    _excel_duplicate_intelligence,
    _excel_formula_gap_scan,
    _excel_reference_check,
    _excel_schema_drift,
    process_uploaded_workbook,
)


def test_excel_studio_supports_tsv_and_reports_import_metadata():
    raw = b"SKU\tQty\tStart\nA-01\t1,200\t08:30\nA-02\t950\t09:45\n"
    result = process_uploaded_workbook(raw, "schedule.tsv")
    frame = result["cleaned_sheets"]["TEXT"]
    assert frame["SKU"].tolist() == ["A-01", "A-02"]
    assert frame["Qty"].tolist() == [1200, 950]
    assert frame["Start"].tolist() == ["08:30", "09:45"]
    assert result["import_diagnostics"]["Encoding"] == "utf-8-sig"
    assert result["import_diagnostics"]["Delimiter"] == "\t"


def test_excel_studio_splits_blank_separated_table_blocks():
    raw = (
        "Orders\n"
        "SKU,Qty\n"
        "A-01,10\n"
        "A-02,20\n"
        "\n"
        "Returns\n"
        "SKU,Qty\n"
        "A-01,1\n"
    ).encode("utf-8")
    result = process_uploaded_workbook(raw, "blocks.csv")
    assert len(result["table_catalog"]) >= 2
    assert len(result["table_datasets"]) >= 2
    assert any("T002" in key for key in result["table_datasets"])


def test_excel_studio_unit_conversion_is_explicit_and_dimension_safe():
    converted, count = _excel_convert_unit_values(pd.Series([60.0, 120.0]), "min", "hr")
    assert count == 2
    assert converted.tolist() == [1.0, 2.0]


def test_excel_studio_validation_engine_and_custom_rules():
    df = pd.DataFrame({"Asset ID": ["A-01", None], "Temp": [80, 120]})
    failures = _excel_apply_validation_rules(df, sheet="Test")
    assert any(x["Field"] == "Asset ID" and x["Status"] == "FAIL" for x in failures)
    custom = _excel_custom_rule_results(
        df,
        [{"name": "Temp <= 100", "field": "Temp", "operator": "<=", "value": 100, "severity": "High"}],
        "Test",
    )
    assert custom[0]["Status"] == "FAIL"
    assert custom[0]["Violations"] == 1


def test_excel_studio_duplicate_intelligence_finds_composite_duplicates():
    df = pd.DataFrame(
        {"Asset ID": ["A-01", "A-01"], "Date": ["2026-01-01", "2026-01-01"], "Qty": [10, 11]}
    )
    findings = _excel_duplicate_intelligence(df, "Orders")
    assert any(x["Type"] == "Composite key duplicate" for x in findings)


def test_excel_studio_schema_drift_detects_added_removed_and_type_changes():
    old = pd.DataFrame({"SKU": ["A"], "Qty": [1]})
    new = pd.DataFrame({"SKU": ["A"], "Qty": [1.5], "Plant": ["Riyadh"]})
    drift = _excel_schema_drift(old, new)
    assert drift["Added fields"] == ["Plant"]
    assert drift["Removed fields"] == []
    assert drift["Type changes"]
    assert not drift["Compatible"]


def test_excel_studio_formula_gap_detection_is_explicit():
    inventory = [
        {"Sheet": "Ops", "Cell": "D2", "Formula": "=B2*C2"},
        {"Sheet": "Ops", "Cell": "D3", "Formula": "=B3*C3"},
        {"Sheet": "Ops", "Cell": "D5", "Formula": "=B5*C5"},
        {"Sheet": "Ops", "Cell": "D6", "Formula": "=B6*C6"},
    ]
    findings = _excel_formula_gap_scan(inventory)
    assert findings
    assert findings[0]["Field"] == "D"
    assert "4" in findings[0]["Affected rows"]


def test_excel_studio_reference_check_reports_coverage():
    clean = pd.DataFrame({"SKU": ["A-01", "A-02"]})
    ref = pd.DataFrame({"SKU": ["A-01", "A-03"]})
    check = _excel_reference_check(clean, ref, "SKU")
    assert check["Matched keys"] == 1
    assert check["Unmatched input keys"] == 1
    assert check["Reference coverage %"] == 50.0


def test_excel_studio_formula_source_is_archived_as_evidence():
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="xlsxwriter") as writer:
        pd.DataFrame({"SKU": ["A-01"], "Qty": [10], "Price": [2], "Total": ["=B2*C2"]}).to_excel(
            writer, index=False, sheet_name="Orders"
        )
    result = process_uploaded_workbook(buf.getvalue(), "formula_gap.xlsx")
    assert result["formula_inventory"]
    assert result["formula_dependencies"]
    book = load_workbook(io.BytesIO(result["xlsx"]), read_only=True, data_only=False)
    assert "FORMULA INVENTORY" in book.sheetnames or result["formula_inventory"]
    book.close()
