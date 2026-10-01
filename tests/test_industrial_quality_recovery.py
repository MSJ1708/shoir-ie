from __future__ import annotations

import ast
import io
from pathlib import Path

import pandas as pd


ROOT = Path(__file__).resolve().parents[1]


def test_app_has_safe_post_module_binding() -> None:
    tree = ast.parse((ROOT / "app.py").read_text(encoding="utf-8"))
    names = {
        node.name
        for node in ast.walk(tree)
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))
    }
    assert "_render_post_module_layers" in names


def test_industrial_relationships_and_mes_workflows(tmp_path) -> None:
    from industrial_platform import (
        init_platform_db,
        load_entity_relationships,
        mes_execution_summary,
        mes_record_event,
        mes_transition_work_order,
        upsert_entities,
        upsert_relationships,
    )

    db = tmp_path / "workspace.db"
    init_platform_db(str(db))
    entities = pd.DataFrame(
        {"ID": ["FAC-001", "MCH-001"], "name": ["Facility", "Machine"], "status": ["Active", "Active"]}
    )
    assert upsert_entities(entities.iloc[[0]], "Facility", "ID", str(db)) == 1
    assert upsert_entities(entities.iloc[[1]], "Machine", "ID", str(db)) == 1
    rel = pd.DataFrame(
        {"From": ["Facility:FAC-001"], "Relationship": ["contains"], "To": ["Machine:MCH-001"]}
    )
    assert upsert_relationships(rel, db_path=str(db)) == 1
    assert len(load_entity_relationships(str(db))) == 1

    import sqlite3
    with sqlite3.connect(db) as con:
        con.execute(
            "INSERT INTO mes_work_orders(work_order,product,quantity,due_date,status,machine,operator,updated_at) VALUES(?,?,?,?,?,?,?,?)",
            ("WO-001", "P-100", 100, "2030-01-01", "Released", "M-01", "", "now"),
        )
        con.commit()

    assert mes_transition_work_order("WO-001", "Running", "operator", str(db))
    mes_record_event("WO-001", "GOOD", 40, "production", "operator", str(db))
    summary = mes_execution_summary(str(db))
    row = summary.loc[summary["Work Order"].eq("WO-001")].iloc[0]
    assert float(row["Good Qty"]) == 40
    assert float(row["WIP Qty"]) == 60


def test_excel_export_is_real_xlsx() -> None:
    from shoir_upgrade import build_excel_report

    payload = build_excel_report("Shoir-IE test", [("Data", pd.DataFrame({"Metric": [1, 2], "Value": [3, 4]}))])
    assert isinstance(payload, bytes) and payload
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=False)
    try:
        assert book.sheetnames
        assert "Data" in book.sheetnames
    finally:
        book.close()
