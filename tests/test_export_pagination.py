"""Artifact-level tests for complete, paginated PDF and PowerPoint exports."""
from io import BytesIO
import hashlib
import json
import os
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches
from PIL import Image
from pypdf import PdfReader

from industrial_platform import export_pdf, export_pptx


def _large_report_table(row_count: int = 67) -> pd.DataFrame:
    return pd.DataFrame({
        "Metric": [f"ROW-{i:03d}" for i in range(1, row_count + 1)],
        "Value": [i * 1.25 for i in range(1, row_count + 1)],
    })



def _save_synthetic_evidence(filename: str, payload: bytes, metadata: dict) -> None:
    """Optionally retain parser-verified synthetic artifacts as a CI evidence bundle."""
    target = os.environ.get("SHOIR_EXPORT_EVIDENCE_DIR", "").strip()
    if not target:
        return

    output_dir = Path(target)
    output_dir.mkdir(parents=True, exist_ok=True)
    artifact_path = output_dir / filename
    artifact_path.write_bytes(payload)

    manifest_path = output_dir / "manifest.json"
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        if not isinstance(manifest, dict):
            manifest = {}
    except (OSError, json.JSONDecodeError):
        manifest = {}

    manifest.update({
        "evidence_type": "synthetic_report_export_validation",
        "synthetic_data_only": True,
        "source_commit": os.environ.get("SHOIR_EXPORT_HEAD_SHA") or os.environ.get("GITHUB_SHA", "local-unidentified"),
    })
    manifest[filename] = {
        **metadata,
        "file": filename,
        "bytes": len(payload),
        "sha256": hashlib.sha256(payload).hexdigest(),
    }
    temporary = manifest_path.with_suffix(".json.tmp")
    temporary.write_text(json.dumps(manifest, indent=2, sort_keys=True), encoding="utf-8")
    temporary.replace(manifest_path)

def test_pdf_export_paginates_and_preserves_rows_past_previous_30_row_limit():
    rows = _large_report_table()
    payload = export_pdf("Complete export verification", [("All metrics", rows)])

    assert payload.startswith(b"%PDF-")
    assert len(payload) > 500
    reader = PdfReader(BytesIO(payload))
    assert len(reader.pages) >= 2

    extracted = "\n".join(page.extract_text() or "" for page in reader.pages)
    for metric, value in zip(rows["Metric"], rows["Value"], strict=True):
        assert metric in extracted
        assert str(value) in extracted
    assert "Value" in extracted
    assert "All metrics" in extracted
    _save_synthetic_evidence("complete_export_verification.pdf", payload, {
        "format": "PDF",
        "row_count": len(rows),
        "parsed_page_count": len(reader.pages),
        "expected_first_row": rows["Metric"].iloc[0],
        "expected_last_row": rows["Metric"].iloc[-1],
        "all_rows_verified": True,
    })


def test_powerpoint_export_splits_all_rows_over_continued_table_slides():
    rows = _large_report_table()
    payload = export_pptx("Complete export verification", [("All metrics", rows)])

    assert payload[:2] == b"PK"
    presentation = Presentation(BytesIO(payload))
    table_slides = [
        slide
        for slide in presentation.slides
        if any(getattr(shape, "has_table", False) for shape in slide.shapes)
    ]
    assert len(table_slides) == (len(rows) + 14) // 15

    exported_metrics = []
    exported_values = []
    for index, slide in enumerate(table_slides):
        table_shapes = [shape for shape in slide.shapes if getattr(shape, "has_table", False)]
        assert len(table_shapes) == 1
        table = table_shapes[0].table
        assert table.cell(0, 0).text == "Metric"
        assert table.cell(0, 1).text == "Value"
        exported_metrics.extend(table.cell(row, 0).text for row in range(1, len(table.rows)))
        exported_values.extend(table.cell(row, 1).text for row in range(1, len(table.rows)))
        if index == 0:
            assert slide.shapes.title.text == "All metrics"
        else:
            assert f"continued {index + 1}" in slide.shapes.title.text

    assert exported_metrics == list(rows["Metric"])
    assert exported_values == [str(value) for value in rows["Value"]]
    _save_synthetic_evidence("complete_export_verification.pptx", payload, {
        "format": "PPTX",
        "row_count": len(rows),
        "parsed_table_slide_count": len(table_slides),
        "expected_first_row": rows["Metric"].iloc[0],
        "expected_last_row": rows["Metric"].iloc[-1],
        "all_rows_verified": True,
    })

class _SyntheticFigure:
    def to_image(self, **kwargs) -> bytes:
        """Return a small 16:9 chart image without external chart services."""
        buffer = BytesIO()
        Image.new("RGB", (160, 90), "white").save(buffer, format="PNG")
        return buffer.getvalue()


def test_powerpoint_titles_tables_and_charts_fit_widescreen_slide_bounds():
    rows = pd.DataFrame({"Metric": ["A", "B"], "Value": [1.0, 2.0]})
    payload = export_pptx(
        "Widescreen layout verification",
        [("Metrics", rows)],
        [("Synthetic chart", _SyntheticFigure())],
    )
    presentation = Presentation(BytesIO(payload))

    width = presentation.slide_width
    height = presentation.slide_height
    assert abs((width / height) - (16 / 9)) < 0.002

    for slide in presentation.slides:
        title = slide.shapes.title
        if title is not None:
            assert title.left >= 0
            assert title.left + title.width <= width
            assert title.top >= 0
            assert title.top + title.height <= height

        content_shapes = [
            shape for shape in slide.shapes
            if getattr(shape, "has_table", False) or shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        ]
        for shape in content_shapes:
            assert shape.left >= 0
            assert shape.top >= 0
            assert shape.left + shape.width <= width
            assert shape.top + shape.height <= height
            if title is not None:
                assert shape.top >= title.top + title.height

