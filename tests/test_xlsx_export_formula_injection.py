"""Regression tests for safe XLSX handling of untrusted text."""
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from shoir_upgrade import build_excel_report


def test_formula_and_url_like_values_remain_inert_text_in_xlsx_report():
    formula = '=HYPERLINK("https://attacker.invalid/collect","Open report")'
    url = "https://attacker.invalid/collect"
    formula_header = '=HYPERLINK("https://attacker.invalid/header","Metric")'
    formula_title = '=HYPERLINK("https://attacker.invalid/title","Report")'

    source = pd.DataFrame({formula_header: [formula, url]})
    payload = build_excel_report(
        formula_title,
        [("Metrics", source)],
        audit=[{"Action": formula, "Source": url}],
        function_reference=False,
    )

    assert payload.startswith(b"PK")
    workbook = load_workbook(BytesIO(payload), data_only=False)
    metrics = workbook["Metrics"]

    # The title, header, and values must remain strings, never spreadsheet formulas.
    for cell_ref, expected in (
        ("A1", formula_title),
        ("A4", formula_header),
        ("A5", formula),
        ("A6", url),
    ):
        cell = metrics[cell_ref]
        assert cell.value == expected
        assert cell.data_type != "f"
        assert cell.hyperlink is None

    audit = workbook["Cleaning Audit"]
    assert audit["A2"].value == formula
    assert audit["B2"].value == url
    assert audit["A2"].data_type != "f"
    assert audit["A2"].hyperlink is None
    assert audit["B2"].hyperlink is None



def test_excel_studio_export_keeps_formula_like_input_as_text():
    from shoir_excel_studio import build_ultimate_workbook

    formula = '=HYPERLINK("https://attacker.invalid/studio","Open")'
    source = pd.DataFrame({"Label": [formula, "https://attacker.invalid/studio"]})
    payload = build_ultimate_workbook(
        "Synthetic safety test",
        {"Cleaned Data": source},
        {"Raw Data": source},
        {"Cleaned Data": [{"Action": formula, "Details": "synthetic"}]},
        {"Cleaned Data": {"Rows": 2, "Columns": 1, "Quality score": 100}},
    )

    workbook = load_workbook(BytesIO(payload), data_only=False)
    found = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value == formula:
                    found.append(cell)
    assert found, "Formula-like input should be retained in the export"
    assert all(cell.data_type != "f" and cell.hyperlink is None for cell in found)


def test_industrial_workbook_export_keeps_formula_like_values_as_text():
    from shoir_industrial_workbook import _serialize_workbook

    formula = '=HYPERLINK("https://attacker.invalid/workbook","Open")'
    payload = _serialize_workbook(
        {"Inputs": pd.DataFrame({"Label": [formula, "https://attacker.invalid/workbook"]})},
        formulas={"Inputs": {"B2": "=1+1"}},
    )

    workbook = load_workbook(BytesIO(payload), data_only=False)
    sheet = workbook["Inputs"]
    # This exporter starts the table at row 3: row 3 is the header, row 4 is data.
    assert sheet["A4"].value == formula
    assert sheet["A4"].data_type != "f"
    assert sheet["A4"].hyperlink is None
    assert sheet["A5"].value == "https://attacker.invalid/workbook"
    assert sheet["A5"].data_type != "f"
    assert sheet["A5"].hyperlink is None
    # The calculation formula inventory is metadata, not a user-executable cell.
    formula_sheets = [ws for ws in workbook.worksheets if "FORMULA" in ws.title.upper()]
    for ws in formula_sheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value == "=1+1":
                    assert cell.data_type != "f"


def test_facility_layout_export_keeps_formula_like_labels_as_text():
    from shoir_facility_layout import _export_workbook

    formula = '=HYPERLINK("https://attacker.invalid/facility","Open")'
    departments = pd.DataFrame([
        {"id": "D1", "name": formula, "x": 0.0, "y": 0.0, "area_sqm": 100.0},
        {"id": "D2", "name": "Assembly", "x": 2.0, "y": 1.0, "area_sqm": 50.0},
    ])
    flows = pd.DataFrame([
        {"from_id": "D1", "to_id": "D2", "loads_day": 10, "distance_m": 2.0},
    ])
    slp = pd.DataFrame([
        {"From": "D1", "To": "D2", "REL": "X", "From Department": formula},
    ])

    payload = _export_workbook(departments, flows, slp)
    workbook = load_workbook(BytesIO(payload), data_only=False)
    matching = [
        cell
        for ws in workbook.worksheets
        for row in ws.iter_rows()
        for cell in row
        if cell.value == formula
    ]
    assert matching
    assert all(cell.data_type != "f" and cell.hyperlink is None for cell in matching)


def test_enterprise_workspace_pack_keeps_formula_like_content_as_text():
    from shoir_enterprise_integration import _build_enterprise_workspace_pack

    formula = '=HYPERLINK("https://attacker.invalid/enterprise","Open")'
    url = "https://attacker.invalid/enterprise"
    payload = _build_enterprise_workspace_pack(
        connectors=pd.DataFrame({"Name": [formula, url]}),
        assets=pd.DataFrame({"Name": ["Synthetic asset"]}),
        activity=pd.DataFrame({"Summary": ["Synthetic activity"]}),
        audit_rows=[{"Area": "Workspace", "Value": formula, "State": url}],
    )

    workbook = load_workbook(BytesIO(payload), data_only=False)
    connector_sheet = workbook["Connectors"]
    for cell_ref, expected in (("A2", formula), ("A3", url)):
        cell = connector_sheet[cell_ref]
        assert cell.value == expected
        assert cell.data_type != "f"
        assert cell.hyperlink is None

    audit_sheet = workbook["Audit"]
    for cell_ref, expected in (("B2", formula), ("C2", url)):
        cell = audit_sheet[cell_ref]
        assert cell.value == expected
        assert cell.data_type != "f"
        assert cell.hyperlink is None
