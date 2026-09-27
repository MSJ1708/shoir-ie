import io

import pandas as pd
from openpyxl import load_workbook

from shoir_excel_studio import clean_dataframe, process_uploaded_workbook

def test_clean_dataframe_handles_messy_export_conservatively():
    table = pd.DataFrame([
        [" Alice ", "0012", "12.5%", "2026-01-05"],
        [" Alice ", "0012", "12.5%", "2026-01-05"],
        ["Bob", "0013", "25%", "2026-01-06"],
    ], columns=["Name", "Order ID", "Rate", "Date"])
    cleaned, audit = clean_dataframe(table)
    assert cleaned.columns.tolist() == ["Name", "Order ID", "Rate", "Date"]
    assert len(cleaned) == 2
    assert cleaned.loc[0, "Name"] == "Alice"
    assert cleaned.loc[0, "Order ID"] == "0012"
    assert abs(float(cleaned.loc[0, "Rate"]) - 0.125) < 1e-9
    assert pd.api.types.is_datetime64_any_dtype(cleaned["Date"])
    assert any(item["Action"] == "Remove duplicate rows" for item in audit)

def test_process_uploaded_workbook_assembles_traceable_xlsx():
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="xlsxwriter") as writer:
        pd.DataFrame([
            ["Operations report", None, None],
            [" SKU ", "Demand", "Date"],
            [" A-01 ", "1,200", "2026-02-01"],
            [" A-02 ", "950", "2026-02-02"],
        ]).to_excel(writer, index=False, header=False, sheet_name="Raw Export")
    result = process_uploaded_workbook(buf.getvalue(), "messy_operations.xlsx")
    cleaned = result["cleaned_sheets"]["Raw Export"]
    assert cleaned.loc[0, "SKU"] == "A-01"
    assert int(cleaned.loc[0, "Demand"]) == 1200
    assert pd.api.types.is_datetime64_any_dtype(cleaned["Date"])
    workbook = load_workbook(io.BytesIO(result["xlsx"]), data_only=False)
    expected = {
        "START HERE", "EXECUTIVE SUMMARY", "DATA DICTIONARY",
        "QUALITY CHECKS", "CLEANING AUDIT", "CLEAN - Raw Export",
    }
    assert expected <= set(workbook.sheetnames)
    raw_sheets = [name for name in workbook.sheetnames if name.startswith("RAW - ")]
    assert raw_sheets and workbook[raw_sheets[0]].sheet_state == "hidden"
    assert workbook["START HERE"]["A1"].value.startswith("Shoir-IE")
    assert result["bundle"][:2] == b"PK"
